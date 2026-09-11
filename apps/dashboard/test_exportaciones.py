from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.dashboard.views import (
    COLUMNAS_EXCEL_ALERTAS,
    crear_excel_alertas,
    crear_excel_completo_amid,
)


class ExportacionesExcelCaracterizacionTests(SimpleTestCase):
    AHORA = datetime(2026, 9, 10, 12, 0)
    FECHA_INICIO = datetime(2026, 8, 28, 0, 0)
    FECHA_FIN = datetime(2026, 9, 11, 0, 0)

    ENCABEZADOS_REGISTROS = [
        "ID", "AMID", "Fecha hora", "Fecha descarga", "Fecha estado",
        "BUSID", "OP", "Versión", "Patente", "TD01", "TD04", "Tabla",
        "Ver tabla", "Latitud", "Longitud", "Porcentaje batería",
        "Tiempo vida", "Fecha registro", "Contiene batería", "Contiene GPS",
        "Contiene tiempo vida", "Error obtener batería", "Error obtener GPS",
        "Error obtener tiempo vida",
    ]

    def serializar(self, workbook):
        salida = BytesIO()
        workbook.save(salida)
        workbook.close()
        salida.seek(0)
        return load_workbook(salida)

    def alerta(self, amid=7500001):
        alerta = {clave: None for _, clave in COLUMNAS_EXCEL_ALERTAS}
        alerta.update({
            "nivel_alerta_global": "ALTA",
            "amid": amid,
            "ubicacion_actual": "Zona sintética",
            "ultimo_estatus": self.AHORA,
            "texto_estatus": "Con estatus",
            "nivel_alerta_gps": "OK",
            "gps_total_hoy": 2,
            "gps_cero_hoy": 1,
            "nivel_alerta_bateria": "ADVERTENCIA",
            "bateria_actual": 0,
            "fecha_actualizacion": self.AHORA,
        })
        return alerta

    def registro(self, identificador, fecha, bateria, latitud, longitud):
        return {
            "id": identificador,
            "amid": 7500001,
            "fecha_hora": fecha,
            "fec_descarga": fecha,
            "fec_estado": fecha,
            "busid": "BUS-TEST",
            "op": "OP-TEST",
            "version": "VTEST",
            "patente": None,
            "td01": None,
            "td04": None,
            "tabla": 3,
            "ver_tabla": "1",
            "latitud": latitud,
            "longitud": longitud,
            "porcentaje_bateria": bateria,
            "tiempo_vida": None,
            "fecha_registro": fecha,
            "is_contiene_bateria": bateria is not None,
            "is_contiene_gps": latitud is not None and longitud is not None,
            "is_contiene_tiempo_vida": False,
            "is_error_obtener_bateria": False,
            "is_error_obtener_gps": False,
            "is_error_obtener_tiempo_vida": False,
        }

    def crear_estandar(self, registros, columnas=None, tabla=None):
        columnas = columnas if columnas is not None else ["08:00", "08:30"]
        tabla = tabla if tabla is not None else [{
            "fecha": "10-09-2026",
            "valores": [{"valor": 0}, {"valor": ""}],
        }]
        with patch(
            "apps.dashboard.views.obtener_rango_fechas_panel",
            return_value=(self.FECHA_INICIO, self.FECHA_FIN),
        ), patch(
            "apps.dashboard.views.obtener_registros_completos_oracle",
            return_value=registros,
        ), patch(
            "apps.dashboard.views.obtener_bloques_bateria_oracle",
            return_value=[],
        ), patch(
            "apps.dashboard.views.construir_tabla_bateria",
            return_value=(columnas, tabla),
        ), patch(
            "apps.dashboard.views.obtener_ahora_referencia",
            return_value=self.AHORA,
        ):
            return crear_excel_completo_amid("7500001")

    def test_excel_alertas_conserva_nombre_y_orden_de_columnas(self):
        wb = self.serializar(crear_excel_alertas(
            [self.alerta()], fecha_generacion=self.AHORA
        ))
        ws = wb["Panel de Alertas"]

        self.assertEqual(wb.sheetnames, ["Panel de Alertas"])
        self.assertEqual(
            [celda.value for celda in ws[4]],
            [titulo for titulo, _ in COLUMNAS_EXCEL_ALERTAS],
        )
        self.assertEqual(len(ws[4]), 25)
        self.assertEqual(ws.freeze_panes, "A5")
        self.assertEqual(list(ws.tables), ["TablaPanelAlertas"])

    def test_excel_alertas_vacio_conserva_encabezados_y_no_crea_tabla(self):
        wb = self.serializar(crear_excel_alertas([], fecha_generacion=self.AHORA))
        ws = wb["Panel de Alertas"]

        self.assertEqual(ws.max_row, 4)
        self.assertEqual(
            [celda.value for celda in ws[4]],
            [titulo for titulo, _ in COLUMNAS_EXCEL_ALERTAS],
        )
        self.assertEqual(list(ws.tables), [])

    def test_excel_alertas_escribe_amid_como_texto_y_conserva_propiedades(self):
        wb = self.serializar(crear_excel_alertas(
            [self.alerta(7500001)], fecha_generacion=self.AHORA
        ))
        ws = wb["Panel de Alertas"]

        self.assertEqual(ws["B5"].value, "7500001")
        self.assertEqual(ws["B5"].data_type, "s")
        self.assertEqual(ws["B5"].number_format, "@")
        self.assertEqual(wb.properties.title, "Panel de Alertas")
        self.assertEqual(wb.properties.subject, "Alertas vigentes de validadores activos")
        self.assertIn("A1:Y1", {str(rango) for rango in ws.merged_cells.ranges})
        self.assertIn("A2:Y2", {str(rango) for rango in ws.merged_cells.ranges})

    def test_excel_estandar_conserva_las_tres_hojas_en_orden(self):
        registros = [self.registro(1, self.AHORA, 80, -33.45, -70.66)]
        wb = self.serializar(self.crear_estandar(registros))

        self.assertEqual(
            wb.sheetnames,
            ["Resumen", "Tabla bateria", "Registros completos"],
        )
        self.assertEqual(wb["Tabla bateria"].freeze_panes, "B6")
        self.assertEqual(wb["Registros completos"].freeze_panes, "A2")

    def test_excel_estandar_conserva_columnas_y_tipo_actual_de_amid(self):
        registros = [self.registro(1, self.AHORA, None, None, None)]
        wb = self.serializar(self.crear_estandar(registros))
        ws = wb["Registros completos"]

        self.assertEqual([celda.value for celda in ws[1]], self.ENCABEZADOS_REGISTROS)
        self.assertEqual(len(ws[1]), 24)
        self.assertEqual(ws["B2"].value, 7500001)
        self.assertEqual(ws["B2"].data_type, "n")
        self.assertIsNone(ws["N2"].value)
        self.assertIsNone(ws["P2"].value)

    def test_tabla_excel_conserva_cero_y_deja_ausencia_vacia(self):
        registros = [self.registro(1, self.AHORA, 0, 0, 0)]
        wb = self.serializar(self.crear_estandar(registros))
        ws = wb["Tabla bateria"]

        self.assertEqual([celda.value for celda in ws[5]], ["Fecha", "08:00", "08:30"])
        self.assertEqual(ws["B6"].value, 0)
        self.assertIsNone(ws["C6"].value)

    def test_resumen_calcula_minimo_maximo_ultima_bateria_y_gps_cero(self):
        registros = [
            self.registro(1, datetime(2026, 9, 10, 8, 0), 80, -33.45, -70.66),
            self.registro(2, datetime(2026, 9, 10, 9, 0), 0, 0, 0),
            self.registro(3, datetime(2026, 9, 10, 10, 0), None, None, None),
            self.registro(4, datetime(2026, 9, 10, 11, 0), 50, -33.46, -70.67),
        ]
        wb = self.serializar(self.crear_estandar(registros))
        ws = wb["Resumen"]
        resumen = {
            ws.cell(fila, 1).value: ws.cell(fila, 2).value
            for fila in range(2, 17)
        }

        self.assertEqual(resumen["AMID"], "7500001")
        self.assertEqual(resumen["Total registros completos"], 4)
        self.assertEqual(resumen["Total registros con batería"], 3)
        self.assertEqual(resumen["Batería mínima"], 0)
        self.assertEqual(resumen["Batería máxima"], 80)
        self.assertEqual(resumen["Última batería registrada"], 50)
        self.assertEqual(resumen["Total registros con GPS"], 3)
        self.assertEqual(resumen["Total GPS 0,0"], 1)

    def test_fechas_aware_se_escriben_sin_timezone(self):
        fecha_aware = timezone.make_aware(datetime(2026, 9, 10, 9, 30))
        registros = [self.registro(1, fecha_aware, 75, -33.45, -70.66)]
        wb = self.serializar(self.crear_estandar(registros))

        fechas = [
            celda.value
            for ws in wb.worksheets
            for fila in ws.iter_rows()
            for celda in fila
            if isinstance(celda.value, datetime)
        ]
        self.assertTrue(fechas)
        self.assertTrue(all(fecha.tzinfo is None for fecha in fechas))
        self.assertEqual(wb["Registros completos"]["C2"].value, datetime(2026, 9, 10, 9, 30))

    def test_excel_sin_registros_retorna_none_y_no_consulta_bloques(self):
        with patch(
            "apps.dashboard.views.obtener_rango_fechas_panel",
            return_value=(self.FECHA_INICIO, self.FECHA_FIN),
        ), patch(
            "apps.dashboard.views.obtener_registros_completos_oracle",
            return_value=[],
        ), patch(
            "apps.dashboard.views.obtener_bloques_bateria_oracle",
        ) as mock_bloques:
            resultado = crear_excel_completo_amid("7500001")

        self.assertIsNone(resultado)
        mock_bloques.assert_not_called()

    def test_excel_estandar_conserva_campos_actuales_del_resumen(self):
        registros = [self.registro(1, self.AHORA, 75, -33.45, -70.66)]
        wb = self.serializar(self.crear_estandar(registros))
        ws = wb["Resumen"]

        indicadores = [ws.cell(fila, 1).value for fila in range(2, 17)]
        self.assertEqual(indicadores, [
            "AMID", "Periodo exportado", "Horario tabla batería",
            "Fecha exportación", "Total registros completos",
            "Total registros con batería", "Total registros con GPS",
            "Total GPS 0,0", "Primera fecha/hora registrada",
            "Última fecha/hora registrada", "Batería mínima", "Batería máxima",
            "Última batería registrada", "Última latitud", "Última longitud",
        ])
