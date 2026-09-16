"""Validacion estructural y de capacidad para archivos Version_DB."""

from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile, ZipFile, is_zipfile

from django.conf import settings
from openpyxl import load_workbook

from apps.dashboard.config.version_zp import (
    COLUMNAS_ESTRUCTURALES_REQUERIDAS,
    VERSION_ZP_EXTENSION,
    VERSION_ZP_SHEET_NAME,
    nombre_columna_oracle,
)


COMPONENTES_XLSX_REQUERIDOS = frozenset({
    "[Content_Types].xml",
    "_rels/.rels",
    "xl/workbook.xml",
    "xl/_rels/workbook.xml.rels",
})


class VersionZPValidationError(ValueError):
    """Error controlado de contrato o capacidad del archivo Version_DB."""


@dataclass(frozen=True)
class VersionZPFileInfo:
    filas: int
    columnas: int
    tamano_comprimido: int
    tamano_descomprimido: int
    ratio_compresion: float


def validar_upload_basico(nombre_archivo, tamano_archivo):
    _validar_extension(nombre_archivo)
    _validar_tamano_comprimido(tamano_archivo)


def validar_archivo_version_zp(ruta_archivo):
    ruta = Path(ruta_archivo)
    validar_upload_basico(ruta.name, ruta.stat().st_size)

    tamano_descomprimido, ratio_compresion = _validar_contenedor_xlsx(ruta)
    filas, columnas = _validar_workbook(ruta)

    return VersionZPFileInfo(
        filas=filas,
        columnas=columnas,
        tamano_comprimido=ruta.stat().st_size,
        tamano_descomprimido=tamano_descomprimido,
        ratio_compresion=ratio_compresion,
    )


def _validar_extension(nombre_archivo):
    if Path(nombre_archivo).suffix.lower() != VERSION_ZP_EXTENSION:
        raise VersionZPValidationError(
            "Formato no soportado. Debes utilizar un archivo .xlsx."
        )


def _validar_tamano_comprimido(tamano_archivo):
    if tamano_archivo > settings.VERSION_ZP_MAX_FILE_BYTES:
        raise VersionZPValidationError(
            "El archivo excede el tamaño máximo comprimido permitido."
        )


def _validar_contenedor_xlsx(ruta):
    if not is_zipfile(ruta):
        raise VersionZPValidationError(
            "El archivo no es un contenedor XLSX válido."
        )

    try:
        with ZipFile(ruta) as archivo_zip:
            entradas = archivo_zip.infolist()
            nombres = {entrada.filename for entrada in entradas}
            faltantes = COMPONENTES_XLSX_REQUERIDOS - nombres
            tiene_hoja = any(
                nombre.startswith("xl/worksheets/") and nombre.endswith(".xml")
                for nombre in nombres
            )

            if faltantes or not tiene_hoja:
                raise VersionZPValidationError(
                    "El archivo ZIP no contiene una estructura XLSX válida."
                )

            total_expandido = sum(entrada.file_size for entrada in entradas)
            total_comprimido = sum(entrada.compress_size for entrada in entradas)

            if total_expandido > settings.VERSION_ZP_MAX_UNCOMPRESSED_BYTES:
                raise VersionZPValidationError(
                    "El contenido XLSX excede el tamaño descomprimido permitido."
                )

            if total_expandido and total_comprimido == 0:
                raise VersionZPValidationError(
                    "El XLSX declara una relación de compresión inválida."
                )

            # Ratio agregado: suma de bytes declarados descomprimidos dividida
            # por la suma de bytes comprimidos de todas las entries del ZIP.
            ratio = (
                total_expandido / total_comprimido
                if total_comprimido
                else 0.0
            )
            if ratio > settings.VERSION_ZP_MAX_COMPRESSION_RATIO:
                raise VersionZPValidationError(
                    "El XLSX excede la relación máxima de compresión permitida."
                )

            return total_expandido, ratio
    except BadZipFile as error:
        raise VersionZPValidationError(
            "El archivo XLSX está corrupto o incompleto."
        ) from error


def _validar_workbook(ruta):
    try:
        workbook = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as error:
        raise VersionZPValidationError(
            "El archivo XLSX está corrupto o no es un workbook válido."
        ) from error

    try:
        if VERSION_ZP_SHEET_NAME not in workbook.sheetnames:
            raise VersionZPValidationError(
                "No se encontró la hoja Version_DB en el Excel "
                "(el nombre distingue mayúsculas y espacios)."
            )

        hoja = workbook[VERSION_ZP_SHEET_NAME]
        columnas = hoja.max_column
        filas_dimensionadas = max(hoja.max_row - 1, 0)

        if columnas > settings.VERSION_ZP_MAX_COLUMNS:
            raise VersionZPValidationError(
                "La hoja Version_DB excede el máximo de columnas permitido."
            )

        if filas_dimensionadas > settings.VERSION_ZP_MAX_ROWS:
            raise VersionZPValidationError(
                "La hoja Version_DB excede el máximo de filas permitido."
            )

        encabezados = next(
            hoja.iter_rows(min_row=1, max_row=1, values_only=True),
            (),
        )
        columnas_normalizadas = {
            nombre
            for encabezado in encabezados
            if encabezado is not None
            for nombre in [nombre_columna_oracle(encabezado)]
            if nombre is not None
        }
        faltantes = sorted(
            COLUMNAS_ESTRUCTURALES_REQUERIDAS - columnas_normalizadas
        )
        if faltantes:
            raise VersionZPValidationError(
                f"Faltan columnas requeridas en el Excel: {faltantes}"
            )

        filas = sum(
            1
            for fila in hoja.iter_rows(min_row=2, values_only=True)
            if any(valor is not None for valor in fila)
        )
        if filas == 0:
            raise VersionZPValidationError(
                "La hoja Version_DB no contiene filas de datos."
            )

        return filas, columnas
    finally:
        workbook.close()
